"""
PRODUCT GENERATOR — Creates Excel templates + PDF planners from keywords
Uses: Claude API for concept design, DesignSystem for Excel, reportlab for PDF
"""
import json, os, re, logging, random, time
from pathlib import Path
from datetime import datetime, timedelta

logger = logging.getLogger("template-machine.generator")

def load_config():
    from config_loader import load_config as _load
    return _load()

# ─── CLAUDE API: TEMPLATE CONCEPT ───────────────────────
def generate_concept(keyword, config):
    """Claude API ile template konsepti üretir."""
    import anthropic
    client = anthropic.Anthropic(api_key=config["anthropic_api_key"])

    prompt = f"""You are a digital product designer for Etsy/Gumroad.
Given keyword: "{keyword}"

Create a complete template concept. Respond ONLY in valid JSON (no markdown, no backticks):
{{
  "product_name": "catchy product name",
  "tagline": "compelling one-line description",
  "excel": {{
    "sheets": [
      {{
        "name": "sheet name",
        "purpose": "what this sheet does",
        "columns": ["Column A", "Column B", "Column C"],
        "sample_rows": 8,
        "has_formulas": true,
        "formula_descriptions": ["SUM totals", "percentage calculations"],
        "has_charts": false
      }}
    ],
    "dashboard_kpis": ["KPI 1", "KPI 2", "KPI 3", "KPI 4"],
    "conditional_formatting": ["red for negative", "green for positive"]
  }},
  "pdf": {{
    "title": "printable planner title",
    "sections": ["Section 1", "Section 2", "Section 3"],
    "page_count": 5,
    "style": "minimal or colorful or professional"
  }},
  "etsy_listing": {{
    "title": "SEO optimized title max 140 chars",
    "description": "compelling 200 word product description",
    "tags": ["tag1", "tag2", "tag3", "tag4", "tag5"]
  }},
  "target_audience": "who buys this",
  "category": "template category"
}}"""

    max_retries = 3
    for attempt in range(max_retries):
        try:
            msg = client.messages.create(
                model=config["model"],
                max_tokens=2000,
                messages=[{"role": "user", "content": prompt}]
            )
            text = next(b.text for b in msg.content if b.type == "text").strip()
            text = text.replace("```json", "").replace("```", "").strip()
            concept = json.loads(text)
            logger.info(f"✅ Konsept üretildi: {concept.get('product_name', keyword)}")
            return concept
        except json.JSONDecodeError as e:
            logger.warning(f"Claude JSON parse hatası (deneme {attempt+1}/{max_retries}): {e}")
            if attempt < max_retries - 1:
                time.sleep(2)
                continue
            logger.error("JSON parse başarısız, fallback kullanılıyor")
            return _fallback_concept(keyword)
        except Exception as e:
            logger.warning(f"Claude API hatası (deneme {attempt+1}/{max_retries}): {e}")
            if attempt < max_retries - 1:
                wait = (attempt + 1) * 3
                logger.info(f"  {wait}s sonra tekrar denenecek...")
                time.sleep(wait)
                continue
            logger.error("Claude API başarısız, fallback kullanılıyor")
            return _fallback_concept(keyword)
    return _fallback_concept(keyword)

def _fallback_concept(keyword):
    words = keyword.title()
    return {
        "product_name": f"{words} — Professional Template",
        "tagline": f"Track, manage, and optimize your {keyword}",
        "excel": {
            "sheets": [
                {"name": "Dashboard", "purpose": "Overview and KPIs", "columns": ["Category", "Value", "Target", "Status"], "sample_rows": 5, "has_formulas": True, "formula_descriptions": ["SUM", "AVERAGE", "IF conditions"], "has_charts": False},
                {"name": "Data", "purpose": "Main data entry", "columns": ["Item", "Category", "Amount", "Date", "Status", "Notes"], "sample_rows": 10, "has_formulas": True, "formula_descriptions": ["Totals", "Conditional status"], "has_charts": False},
            ],
            "dashboard_kpis": ["Total", "Average", "Count", "Top Item"],
            "conditional_formatting": ["red below target", "green above target"]
        },
        "pdf": {"title": f"{words} Planner", "sections": ["Overview", "Weekly Plan", "Tracking", "Notes"], "page_count": 4, "style": "professional"},
        "etsy_listing": {
            "title": f"{words} Template Excel Spreadsheet Google Sheets Planner Tracker Digital Download",
            "description": f"Professional {keyword} template to help you stay organized and track everything in one place. Includes automated calculations, dashboard, and printable planner.",
            "tags": keyword.split()[:5] + ["template", "spreadsheet", "planner", "digital download", "tracker"]
        },
        "target_audience": "professionals and small business owners",
        "category": "Templates"
    }

# ─── SAMPLE DATA GENERATOR ────────────────────────────
def _generate_sample_value(col_name, row_index):
    """Kolon adına göre mantıklı örnek veri üretir."""
    col = col_name.lower().strip()

    # Tarih kolonları
    if any(w in col for w in ["date", "tarih", "day", "when", "deadline"]):
        base = datetime.now() - timedelta(days=30)
        return (base + timedelta(days=row_index * 3)).strftime("%Y-%m-%d")

    # Para/fiyat kolonları
    if any(w in col for w in ["amount", "price", "cost", "total", "budget", "revenue", "expense", "tutar", "fiyat", "salary", "payment"]):
        amounts = [120, 250, 89, 340, 175, 420, 95, 560, 210, 310, 150, 480]
        return amounts[row_index % len(amounts)]

    # Yüzde kolonları
    if any(w in col for w in ["percent", "%", "rate", "ratio", "progress", "completion"]):
        pcts = [0.85, 0.72, 0.93, 0.61, 0.78, 0.45, 0.88, 0.67, 0.91, 0.54]
        return pcts[row_index % len(pcts)]

    # Miktar/sayı kolonları
    if any(w in col for w in ["quantity", "qty", "count", "number", "units", "hours", "adet", "miktar"]):
        nums = [5, 12, 3, 8, 15, 7, 20, 4, 10, 6]
        return nums[row_index % len(nums)]

    # Durum kolonları
    if any(w in col for w in ["status", "state", "durum", "progress"]):
        statuses = ["Complete", "In Progress", "Pending", "Complete", "In Progress", "Not Started", "Complete", "Pending", "In Progress", "Complete"]
        return statuses[row_index % len(statuses)]

    # Kategori kolonları
    if any(w in col for w in ["category", "type", "group", "kategori", "tür"]):
        cats = ["Marketing", "Operations", "Finance", "HR", "Sales", "IT", "Legal", "Support", "Product", "Design"]
        return cats[row_index % len(cats)]

    # Öncelik kolonları
    if any(w in col for w in ["priority", "öncelik", "importance"]):
        priorities = ["High", "Medium", "Low", "High", "Medium", "Critical", "Low", "High", "Medium", "Low"]
        return priorities[row_index % len(priorities)]

    # İsim/ad kolonları
    if any(w in col for w in ["name", "item", "task", "project", "ad", "isim", "title", "description"]):
        items = [
            "Q1 Marketing Campaign", "Website Redesign", "Employee Training",
            "Budget Review", "Product Launch", "Client Meeting Prep",
            "Inventory Audit", "Social Media Strategy", "Annual Report",
            "Team Building Event", "Supply Chain Optimization", "Customer Survey"
        ]
        return items[row_index % len(items)]

    # Not kolonları
    if any(w in col for w in ["note", "comment", "remarks", "not", "açıklama"]):
        notes = ["On track", "Needs review", "Updated", "Approved", "Follow up", "Completed ahead of schedule", "Waiting for input", "In review", "Scheduled", "Done"]
        return notes[row_index % len(notes)]

    # Email
    if any(w in col for w in ["email", "e-mail", "mail"]):
        names = ["alice", "bob", "carol", "dave", "eve", "frank", "grace", "henry", "iris", "jack"]
        return f"{names[row_index % len(names)]}@example.com"

    # Varsayılan: genel metin
    defaults = ["Item A", "Item B", "Item C", "Item D", "Item E", "Item F", "Item G", "Item H", "Item I", "Item J"]
    return defaults[row_index % len(defaults)]


def _is_numeric_column(col_name):
    """Kolonun sayısal veri içerip içermediğini kontrol eder."""
    col = col_name.lower().strip()
    return any(w in col for w in [
        "amount", "price", "cost", "total", "budget", "revenue", "expense",
        "tutar", "fiyat", "salary", "payment", "quantity", "qty", "count",
        "number", "units", "hours", "adet", "miktar", "percent", "%", "rate",
        "ratio", "progress", "completion"
    ])


# ─── EXCEL GENERATION ───────────────────────────────────
def generate_excel(keyword, concept, config):
    """Konsepte göre premium Excel template üretir (DesignSystem ile)."""
    from design_system import create_styled_workbook, THEMES
    from openpyxl.utils import get_column_letter

    # Tema seçimi
    themes_list = list(THEMES.keys())
    if config["design"]["theme_rotation"]:
        theme_name = random.choice(themes_list)
    else:
        theme_name = themes_list[0]

    wb, ds = create_styled_workbook(theme=theme_name)

    sheets_config = concept.get("excel", {}).get("sheets", [])
    kpis = concept.get("excel", {}).get("dashboard_kpis", ["Total", "Average", "Count", "Status"])

    for si, sheet_cfg in enumerate(sheets_config):
        ws = wb.active if si == 0 else wb.create_sheet()
        raw_title = sheet_cfg.get("name", f"Sheet{si+1}")
        # Excel sheet names: max 31 chars, no \ / ? * [ ] :
        ws.title = re.sub(r'[\\/?*\[\]:]', '-', str(raw_title))[:31] or f"Sheet{si+1}"
        columns = sheet_cfg.get("columns", ["A", "B", "C", "D"])
        num_cols = len(columns)
        sample_rows = sheet_cfg.get("sample_rows", 8)

        # ── Title banner ──
        title_text = concept.get("product_name", keyword.title()) if si == 0 else sheet_cfg.get("name", "")
        subtitle_text = concept.get("tagline", "") if si == 0 else sheet_cfg.get("purpose", "")
        merge_end = get_column_letter(num_cols + 1)
        ds.style_dashboard_title(ws, title_text, subtitle_text, f"A1:{merge_end}1", row=1)

        # Spacer row
        ws.row_dimensions[3].height = 6
        for c in range(1, num_cols + 2):
            ws.cell(row=3, column=c).fill = ds.fill_bg_light

        # ── KPI cards for dashboard (first sheet) ──
        if si == 0 and kpis:
            # Find numeric columns for KPI formulas
            numeric_cols = [i for i, col in enumerate(columns) if _is_numeric_column(col)]
            header_start = 4 + 3  # KPI cards take 3 rows (4,5,6) → header at 7

            for ki, kpi_label in enumerate(kpis[:4]):
                col = ki + 1
                # Use formula if we have numeric columns, else placeholder
                if numeric_cols and ki < len(numeric_cols):
                    nc = numeric_cols[ki % len(numeric_cols)]
                    data_col_letter = get_column_letter(nc + 1)
                    data_start = header_start + 1
                    data_end = header_start + sample_rows
                    if "total" in kpi_label.lower() or "sum" in kpi_label.lower():
                        formula = f"=SUM({data_col_letter}{data_start}:{data_col_letter}{data_end})"
                    elif "average" in kpi_label.lower() or "avg" in kpi_label.lower():
                        formula = f"=AVERAGE({data_col_letter}{data_start}:{data_col_letter}{data_end})"
                    elif "count" in kpi_label.lower():
                        formula = f"=COUNTA({data_col_letter}{data_start}:{data_col_letter}{data_end})"
                    elif "max" in kpi_label.lower() or "top" in kpi_label.lower() or "high" in kpi_label.lower():
                        formula = f"=MAX({data_col_letter}{data_start}:{data_col_letter}{data_end})"
                    elif "min" in kpi_label.lower() or "low" in kpi_label.lower():
                        formula = f"=MIN({data_col_letter}{data_start}:{data_col_letter}{data_end})"
                    else:
                        formula = f"=SUM({data_col_letter}{data_start}:{data_col_letter}{data_end})"
                    ds.style_kpi_card(ws, row=4, col=col, label=kpi_label, value_cell_ref=formula)
                else:
                    ds.style_kpi_card(ws, row=4, col=col, label=kpi_label, value_cell_ref="—")

        else:
            header_start = 4

        # ── Header row ──
        ds.style_header_row(ws, row=header_start, col_start=1, col_end=num_cols)
        for ci, col_name in enumerate(columns):
            ws.cell(row=header_start, column=ci + 1, value=col_name)

        # ── Data rows with sample data ──
        for ri in range(sample_rows):
            r = header_start + 1 + ri
            for ci, col_name in enumerate(columns):
                ws.cell(row=r, column=ci + 1, value=_generate_sample_value(col_name, ri))

        data_start_row = header_start + 1
        data_end_row = header_start + sample_rows
        ds.style_data_rows(ws, data_start_row, data_end_row, col_start=1, col_end=num_cols)

        # ── Formulas: Total row for numeric columns ──
        if sheet_cfg.get("has_formulas", False):
            total_row = data_end_row + 1
            ws.cell(row=total_row, column=1, value="TOTAL")
            for ci, col_name in enumerate(columns):
                if _is_numeric_column(col_name):
                    col_letter = get_column_letter(ci + 1)
                    ws.cell(row=total_row, column=ci + 1,
                            value=f"=SUM({col_letter}{data_start_row}:{col_letter}{data_end_row})")
            ds.style_total_row(ws, total_row, col_start=1, col_end=num_cols)

            # Conditional formatting on numeric columns
            for ci, col_name in enumerate(columns):
                if _is_numeric_column(col_name):
                    col_letter = get_column_letter(ci + 1)
                    ds.add_data_bars(ws, col_letter, data_start_row, data_end_row)

        # ── Column widths ──
        widths = {}
        for ci in range(num_cols):
            letter = get_column_letter(ci + 1)
            col_name = columns[ci].lower()
            if any(w in col_name for w in ["name", "item", "task", "project", "description", "title", "note"]):
                widths[letter] = 28
            elif any(w in col_name for w in ["email", "mail"]):
                widths[letter] = 24
            else:
                widths[letter] = 16
        ds.auto_column_widths(ws, widths)

        # ── Finalize ──
        ds.finalize_sheet(ws, freeze=f"A{header_start + 1}", gridlines=False)

    # Save
    output_dir = Path(config["output_dir"])
    output_dir.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r'[^\w\s-]', '', keyword).strip().replace(' ', '_')
    filename = f"{safe_name}_template.xlsx"
    filepath = output_dir / filename
    wb.save(str(filepath))
    logger.info(f"📊 Excel oluşturuldu: {filepath} (tema: {theme_name})")
    return str(filepath), theme_name


# ─── PDF THEME MAPPING ────────────────────────────────
PDF_THEME_COLORS = {
    "ocean":    {"accent": "#1B4965", "dark": "#0B2545", "light": "#EEF4F8", "border": "#D1DCE8", "muted": "#6B7B8D"},
    "sunset":   {"accent": "#6C2E7A", "dark": "#2D1B3D", "light": "#FDF5F0", "border": "#E8D5E0", "muted": "#8B7B8D"},
    "forest":   {"accent": "#1A472A", "dark": "#0D2818", "light": "#F0F7F2", "border": "#C8DCC8", "muted": "#6B8B6D"},
    "midnight": {"accent": "#1A1A2E", "dark": "#0A0A1A", "light": "#F2F2F8", "border": "#D0D0E0", "muted": "#6B6B8D"},
}

# ─── PDF GENERATION ─────────────────────────────────────
def generate_pdf(keyword, concept, config, theme_name="ocean"):
    """Konsepte göre printable PDF planner üretir (tema uyumlu)."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
        from reportlab.lib.enums import TA_CENTER
    except ImportError:
        logger.warning("reportlab not installed, skipping PDF generation")
        return None

    output_dir = Path(config["output_dir"])
    output_dir.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r'[^\w\s-]', '', keyword).strip().replace(' ', '_')
    filename = f"{safe_name}_planner.pdf"
    filepath = output_dir / filename

    pdf_cfg = concept.get("pdf", {})
    sections = pdf_cfg.get("sections", ["Overview", "Planning", "Tracking", "Notes"])
    title = pdf_cfg.get("title", f"{keyword.title()} Planner")

    # Tema renklerini al
    tc = PDF_THEME_COLORS.get(theme_name, PDF_THEME_COLORS["ocean"])
    ACCENT = colors.HexColor(tc["accent"])
    DARK = colors.HexColor(tc["dark"])
    SILK = colors.HexColor(tc["border"])
    MUTED = colors.HexColor(tc["muted"])

    doc = SimpleDocTemplate(str(filepath), pagesize=A4, topMargin=20*mm, bottomMargin=20*mm, leftMargin=20*mm, rightMargin=20*mm)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle("CustomTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=28, textColor=DARK, spaceAfter=6)
    subtitle_style = ParagraphStyle("Subtitle", parent=styles["Normal"], fontName="Helvetica", fontSize=11, textColor=MUTED, spaceAfter=30)
    section_style = ParagraphStyle("Section", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=16, textColor=ACCENT, spaceBefore=20, spaceAfter=10)
    body_style = ParagraphStyle("Body", parent=styles["Normal"], fontName="Helvetica", fontSize=10, textColor=DARK, spaceAfter=8, leading=14)

    elements = []

    # Cover page
    elements.append(Spacer(1, 80))
    elements.append(Paragraph(title, title_style))
    elements.append(Paragraph(concept.get("tagline", f"Your complete {keyword} planning toolkit"), subtitle_style))

    # Accent line
    line_data = [["" for _ in range(6)]]
    line_table = Table(line_data, colWidths=[28*mm]*6, rowHeights=[2*mm])
    line_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), ACCENT)]))
    elements.append(line_table)
    elements.append(Spacer(1, 40))

    info_text = f"""
    <b>What's Inside:</b><br/>
    {'<br/>'.join(f'&#10003; {s}' for s in sections)}<br/><br/>
    <b>How to Use:</b><br/>
    Print this planner and fill it in daily/weekly.<br/>
    Use alongside the Excel template for digital tracking.
    """
    elements.append(Paragraph(info_text, body_style))
    elements.append(Spacer(1, 20))

    footer_text = f"&copy; {datetime.now().year} {config['design']['brand_name']} &mdash; {config['design']['tagline']}"
    footer_style = ParagraphStyle("Footer", parent=styles["Normal"], fontName="Helvetica", fontSize=8, textColor=MUTED, alignment=TA_CENTER)
    elements.append(Paragraph(footer_text, footer_style))
    elements.append(PageBreak())

    # Section pages
    for section in sections:
        elements.append(Paragraph(section, section_style))

        sl = Table([[""] * 4], colWidths=[42*mm]*4, rowHeights=[1.5*mm])
        sl.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), ACCENT)]))
        elements.append(sl)
        elements.append(Spacer(1, 15))

        for _ in range(16):
            row = Table([[""] * 2], colWidths=[85*mm, 85*mm], rowHeights=[12*mm])
            row.setStyle(TableStyle([
                ("LINEBELOW", (0, 0), (-1, -1), 0.5, SILK),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]))
            elements.append(row)

        elements.append(PageBreak())

    # Notes page
    elements.append(Paragraph("Notes", section_style))
    sl = Table([[""] * 4], colWidths=[42*mm]*4, rowHeights=[1.5*mm])
    sl.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), ACCENT)]))
    elements.append(sl)
    elements.append(Spacer(1, 15))
    for _ in range(22):
        row = Table([[""] * 1], colWidths=[170*mm], rowHeights=[10*mm])
        row.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.5, SILK)]))
        elements.append(row)

    doc.build(elements)
    logger.info(f"📄 PDF oluşturuldu: {filepath} (tema: {theme_name})")
    return str(filepath)


# ─── MAIN ───────────────────────────────────────────────

def generate_products(keyword, config=None):
    """Bir keyword için tüm ürünleri üretir."""
    if config is None:
        config = load_config()

    logger.info(f"🏭 Ürün üretimi başlıyor: '{keyword}'")

    concept = generate_concept(keyword, config)
    excel_path, theme = generate_excel(keyword, concept, config)
    pdf_path = generate_pdf(keyword, concept, config, theme_name=theme)

    # Preview görseli
    from preview_generator import generate_preview
    preview_path = generate_preview(keyword, concept, theme_name=theme, output_dir=config["output_dir"])

    # Fiyatlandırma
    excel_price = config["pricing"]["excel_base"]
    pdf_price = config["pricing"]["pdf_base"]
    bundle_price = round((excel_price + pdf_price) * (1 - config["pricing"]["bundle_discount"]), 2)

    return {
        "keyword": keyword,
        "concept": concept,
        "theme": theme,
        "files": {
            "excel": excel_path,
            "pdf": pdf_path,
            "preview": preview_path,
        },
        "pricing": {
            "excel": excel_price,
            "pdf": pdf_price,
            "bundle": bundle_price,
        },
        "listing": concept.get("etsy_listing", {}),
    }

if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(message)s")
    result = generate_products("meal prep planner")
    print(json.dumps({k: v for k, v in result.items() if k != "concept"}, indent=2))
