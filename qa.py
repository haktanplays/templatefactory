"""
QA GATE — Hiçbir hatalı ürün yayına çıkmasın
=============================================
Her ürün (Excel + PDF + görsel + listing) satışa çıkmadan önce buradan geçer.
Excel: LibreOffice ile yeniden hesaplanır, hata değeri (#VALUE! #REF! #NAME? ...)
aranır, KPI/metrik/insight hücreleri, TOTAL satırları, dropdown'lar, sayfa adları,
yer tutucu metinler ve birleşik hücreler kontrol edilir.
PDF: sayfa sayısı, her sayfanın başlığı, boş sayfa. Görsel: boyut ve içerik.
Listing: Etsy sınırları ve yer tutucu metin.
"""
import json, logging, re, shutil, subprocess, tempfile, warnings
from datetime import datetime, date
from pathlib import Path

logger = logging.getLogger("template-machine.qa")

ERROR_VALUES = ("#VALUE!", "#REF!", "#NAME?", "#DIV/0!", "#N/A", "#NUM!", "#NULL!")
PLACEHOLDER_WORDS = re.compile(r"\b(lorem|ipsum|placeholder|tbd|todo|xxx+|item [a-j]\b|sample text|your text here|insert here)\b", re.I)
MIN_SAMPLE_ROWS = 5


class QAReport:
    def __init__(self):
        self.checks = []

    def add(self, name, ok, detail=""):
        self.checks.append({"name": name, "ok": bool(ok), "detail": detail})
        (logger.info if ok else logger.warning)(f"  {'✓' if ok else '✗'} {name}{': ' + detail if detail and not ok else ''}")
        return ok

    @property
    def passed(self):
        return all(c["ok"] for c in self.checks)

    @property
    def failures(self):
        return [f"{c['name']}: {c['detail']}" for c in self.checks if not c["ok"]]

    def to_dict(self):
        return {"passed": self.passed, "checked_at": datetime.now().isoformat(), "checks": self.checks}


# ─── LIBREOFFICE RECALC ─────────────────────────────────
def recalculate(xlsx_path, timeout=180):
    """Dosyayı LibreOffice'e açtırıp hesaplanmış değerlerle kaydettirir; yeni dosya yolunu döndürür (araç yoksa None)."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    tmp = Path(tempfile.mkdtemp(prefix="tm_qa_"))
    try:
        subprocess.run([soffice, "--headless", "--convert-to", "xlsx", "--outdir", str(tmp), str(xlsx_path)],
                       check=True, capture_output=True, timeout=timeout)
        out = tmp / Path(xlsx_path).name
        return str(out) if out.exists() else None
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as e:
        logger.warning(f"LibreOffice yeniden hesaplama başarısız: {e}")
        return None


# ─── EXCEL ──────────────────────────────────────────────
def check_excel(xlsx_path, spec, report, strict=True):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    from excel_builder import HEADER_ROW

    try:
        wb = load_workbook(xlsx_path)
    except Exception as e:
        report.add("excel.opens", False, f"openpyxl açamadı: {e}")
        return
    report.add("excel.opens", True)

    names = wb.sheetnames
    bad = [n for n in names if len(n) > 31 or re.search(r'[\\/?*\[\]:]', n)]
    report.add("excel.sheet_names", not bad and len(set(n.lower() for n in names)) == len(names), f"geçersiz/çift: {bad}")
    expected = ["Start Here", "Dashboard"] + [s["name"] for s in spec["excel"]["sheets"]]
    report.add("excel.sheet_order", names == expected, f"{names} != {expected}")

    # Dropdown doğrulamaları: 255 karakter sınırı ve boş liste
    dv_problems = []
    for ws in wb.worksheets:
        for dv in ws.data_validations.dataValidation:
            if dv.type == "list":
                if not dv.formula1 or len(dv.formula1) > 255:
                    dv_problems.append(f"{ws.title}:{dv.sqref} ({len(dv.formula1 or '')} kr)")
    report.add("excel.dropdowns", not dv_problems, "; ".join(dv_problems))

    # Birleşik hücreler çakışmasın
    overlap = []
    for ws in wb.worksheets:
        ranges = list(ws.merged_cells.ranges)
        for i, a in enumerate(ranges):
            for b in ranges[i + 1:]:
                if a.min_row <= b.max_row and b.min_row <= a.max_row and a.min_col <= b.max_col and b.min_col <= a.max_col:
                    overlap.append(f"{ws.title}:{a}&{b}")
    report.add("excel.merges", not overlap, "; ".join(overlap[:5]))

    # Veri sayfaları: başlık, örnek satır sayısı, yer tutucu metin
    for sh in spec["excel"]["sheets"]:
        ws = wb[sh["name"]]
        n = len(sh["columns"])
        headers = [ws.cell(row=HEADER_ROW, column=c).value for c in range(1, n + 1)]
        report.add(f"excel.{sh['name']}.headers", all(headers) and headers == [c["name"] for c in sh["columns"]], f"{headers}")
        min_rows = 7 if any("week" in sh["name"].lower() for _ in [0]) else MIN_SAMPLE_ROWS
        report.add(f"excel.{sh['name']}.sample_rows", len(sh["rows"]) >= min(min_rows, 7), f"{len(sh['rows'])} satır (<{min_rows})")
        texts = " ".join(str(v) for row in sh["rows"] for v in row if isinstance(v, str))
        m = PLACEHOLDER_WORDS.search(texts)
        report.add(f"excel.{sh['name']}.no_placeholders", not m, f"yer tutucu: '{m.group(0)}'" if m else "")
        report.add(f"excel.{sh['name']}.freeze", ws.freeze_panes == f"A{HEADER_ROW + 1}", str(ws.freeze_panes))
        if sh.get("chart"):
            report.add(f"excel.{sh['name']}.chart", len(ws._charts) == 1, f"{len(ws._charts)} grafik")

    dash = wb["Dashboard"]
    n_kpi = len(spec["excel"]["dashboard"]["kpis"])
    report.add("excel.dashboard.kpi_count", n_kpi == 4, f"{n_kpi} KPI (4 bekleniyor)")
    report.add("excel.dashboard.charts", len(dash._charts) == len(spec["excel"]["dashboard"]["charts"]) and len(dash._charts) >= 1,
               f"{len(dash._charts)} grafik")

    # ── Formül değerlendirme (LibreOffice) ──
    calc_path = recalculate(xlsx_path)
    if not calc_path:
        report.add("excel.formulas_evaluated", not strict,
                   "LibreOffice bulunamadı: formüller hesaplatılamadı. Kurulum: apt install libreoffice-calc / brew install --cask libreoffice")
        return
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")      # LibreOffice'in eklediği uzantılar openpyxl'i ilgilendirmez
        wbv = load_workbook(calc_path, data_only=True)

    errors = []
    for ws in wbv.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and (v in ERROR_VALUES or v.startswith("Err:")):
                    errors.append(f"{ws.title}!{cell.coordinate}={v}")
    report.add("excel.no_formula_errors", not errors, "; ".join(errors[:8]) + (f" (+{len(errors)-8})" if len(errors) > 8 else ""))

    # KPI + metrik değerleri sayısal, insight'lar dolu metin
    dv = wbv["Dashboard"]
    layout = _dashboard_layout(dash)
    for kind, cells in (("kpi", layout["kpi_values"]), ("metric", layout["metric_values"])):
        vals = [dv[c].value for c in cells]
        ok = all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in vals)
        report.add(f"excel.dashboard.{kind}_values", ok, f"{list(zip(cells, vals))}")
    ins = [dv[c].value for c in layout["insight_cells"]]
    report.add("excel.dashboard.insights", all(isinstance(v, str) and len(v.strip()) > 10 and v not in ERROR_VALUES for v in ins) if ins else True,
               f"{ins}")

    # TOTAL satırları: hesaplanan toplam, örnek verinin Python toplamına eşit mi
    for sh in spec["excel"]["sheets"]:
        ws_f, ws_v = wb[sh["name"]], wbv[sh["name"]]
        data_cols = [c for c in sh["columns"] if not c.get("formula")]
        total_row = None
        for r in range(HEADER_ROW + 1, ws_f.max_row + 1):
            if ws_f.cell(row=r, column=1).value == "TOTAL":
                total_row = r
                break
        for ci, col in enumerate(sh["columns"]):
            if col.get("total") == "sum" and not col.get("formula") and total_row:
                di = data_cols.index(col)
                expected = sum(float(row[di]) for row in sh["rows"] if isinstance(row[di], (int, float)))
                got = ws_v.cell(row=total_row, column=ci + 1).value
                ok = isinstance(got, (int, float)) and abs(got - expected) < 0.01
                report.add(f"excel.{sh['name']}.total.{col['name']}", ok, f"hesaplanan {got}, beklenen {expected}")
            elif col.get("formula") and sh["rows"]:
                # formül kolonu örnek satırlarda sayısal sonuç vermeli
                v = ws_v.cell(row=HEADER_ROW + 1, column=ci + 1).value
                report.add(f"excel.{sh['name']}.formula.{col['name']}", isinstance(v, (int, float)), f"ilk satır değeri: {v!r}")


def _dashboard_layout(dash_ws):
    """Dashboard'daki KPI/metrik değer hücrelerini ve insight hücrelerini bulur (excel_builder düzeni)."""
    from excel_builder import DASH
    kpi_values = [f"{col}{DASH['kpi_row'] + 1}" for col in "ABCD" if dash_ws[f"{col}{DASH['kpi_row']}"].value]
    metric_values = [f"{col}{DASH['metric_row'] + 1}" for col in "ABCD" if dash_ws[f"{col}{DASH['metric_row']}"].value]
    insight_cells = []
    r = DASH["insight_row"] + 1
    while dash_ws[f"A{r}"].value:
        insight_cells.append(f"A{r}")
        r += 1
    return {"kpi_values": kpi_values, "metric_values": metric_values, "insight_cells": insight_cells}


# ─── PDF ────────────────────────────────────────────────
def check_pdf(pdf_path, spec, report):
    try:
        from pypdf import PdfReader
        reader = PdfReader(pdf_path)
        n = len(reader.pages)
    except Exception as e:
        report.add("pdf.opens", False, str(e))
        return
    report.add("pdf.opens", True)
    pages = spec["pdf"]["pages"]
    report.add("pdf.page_count", len(pages) <= n <= len(pages) + 6, f"{n} sayfa, spec {len(pages)} sayfa tipi")
    texts = [(p.extract_text() or "") for p in reader.pages]
    blank = [i + 1 for i, t in enumerate(texts) if len(t.strip()) < 5]
    report.add("pdf.no_blank_pages", not blank, f"boş sayfalar: {blank}")
    missing = []
    alltext = " ".join(texts).lower()
    for p in pages:
        title = p.get("title")
        if title and title.lower()[:30] not in alltext:
            missing.append(title)
    report.add("pdf.titles_rendered", not missing, f"eksik: {missing}")
    report.add("pdf.cover_title", spec["pdf"].get("title", "").lower()[:20] in texts[0].lower(), texts[0][:60])
    m = PLACEHOLDER_WORDS.search(alltext)
    report.add("pdf.no_placeholders", not m, m.group(0) if m else "")


# ─── GÖRSEL ─────────────────────────────────────────────
def check_preview(png_path, report):
    try:
        from PIL import Image, ImageStat
        img = Image.open(png_path).convert("RGB")
    except Exception as e:
        report.add("preview.opens", False, str(e))
        return
    report.add("preview.opens", True)
    report.add("preview.size", img.size == (2000, 1600), str(img.size))
    stat = ImageStat.Stat(img)
    report.add("preview.has_content", sum(stat.stddev) / 3 > 40, f"stddev {sum(stat.stddev)/3:.1f}")
    # sağ yarıda ekran görüntüsü (beyaz alan) var mı
    pixels = list(img.crop((1000, 200, 2000, 1400)).getdata())[::50]
    white = sum(1 for px in pixels if min(px) > 235)
    report.add("preview.screenshot_present", white > len(pixels) * 0.15, f"beyaz oran {white/len(pixels):.2f}")


# ─── LISTING ────────────────────────────────────────────
def check_listing(listing, report):
    title = listing.get("title", "")
    report.add("listing.title", 20 <= len(title) <= 140, f"{len(title)} karakter")
    tags = listing.get("tags", [])
    report.add("listing.tags", len(tags) == 13 and all(1 <= len(t) <= 20 for t in tags) and len(set(tags)) == 13, f"{len(tags)} etiket")
    words = len(listing.get("description", "").split())
    report.add("listing.description", words >= 80, f"{words} kelime (<80)")
    m = PLACEHOLDER_WORDS.search(title + " " + listing.get("description", ""))
    report.add("listing.no_placeholders", not m, m.group(0) if m else "")


# ─── GİRİŞ NOKTASI ──────────────────────────────────────
def run_qa(product, config):
    """Ürün sözlüğünü (generate_products çıktısı) kontrol eder; QAReport döndürür ve yanına _qa.json yazar."""
    strict = config.get("qa", {}).get("strict", True)
    report = QAReport()
    spec = product["concept"]
    files = product["files"]
    logger.info(f"🔎 QA: {product['keyword']}")
    check_excel(files["excel"], spec, report, strict=strict)
    check_pdf(files["pdf"], spec, report)
    check_preview(files["preview"], report)
    check_listing(product.get("listing", {}), report)
    qa_path = Path(files["excel"]).with_name(Path(files["excel"]).stem.removesuffix("_template") + "_qa.json")
    with open(qa_path, "w", encoding="utf-8") as f:
        json.dump(report.to_dict(), f, indent=2, ensure_ascii=False)
    n_fail = len(report.failures)
    logger.info(f"  {'✅ QA geçti' if report.passed else f'❌ QA {n_fail} hata'} ({len(report.checks)} kontrol) → {qa_path}")
    return report
